"""Экспорт утверждённых данных в файл по шаблону пользователя.

Поддерживаемые шаблоны:
  * XML (и любой текстовый формат: .txt, .csv) — синтаксис Jinja2:
        <Importer_name>{{ header.importer.name }}</Importer_name>
        {% for item in items %}<Item>...{{ item.hs_code }}...</Item>{% endfor %}
  * Excel (.xlsx) — плейсхолдеры в ячейках:
        {{ header.importer.name }}            — значение из шапки
        {{ item.description }}, {{ item.total_value }} — строка-образец товара:
        строка, где есть «item.», размножается по числу позиций.
  * Без шаблона — встроенный XML (app/export_templates/default_declaration.xml).

Шаблоны загружает пользователь, поэтому они исполняются в SandboxedEnvironment:
из шаблона нельзя добраться до файлов сервера или внутренних объектов Python.
"""

import io
import re
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any

from jinja2 import TemplateError
from jinja2.sandbox import SandboxedEnvironment
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.i18n import t
from app.schemas import DeclarationData

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "export_templates" / "default_declaration.xml"
TEXT_TEMPLATE_TYPES = {".xml": "application/xml", ".txt": "text/plain", ".csv": "text/csv"}
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Ячейка целиком состоит из одного {{ выражения }} (внутри нет других «{{» / «}}»)
_SINGLE_EXPRESSION = re.compile(r"^\s*\{\{((?:(?!\{\{|\}\}).)+)\}\}\s*$", re.DOTALL)


class TemplateRenderError(Exception):
    pass


def _num(value: Any, digits: int = 2) -> str:
    """Фильтр {{ x|num(3) }}: число с фиксированным числом знаков, точка как разделитель."""
    if value is None or value == "":
        return ""
    return f"{float(value):.{digits}f}"


def _make_env(autoescape: bool) -> SandboxedEnvironment:
    env = SandboxedEnvironment(
        autoescape=autoescape,
        keep_trailing_newline=True,
        # None выводим как пустую строку, а не «None»
        finalize=lambda value: "" if value is None else value,
    )
    env.filters["num"] = _num
    return env


def _context(data: DeclarationData) -> dict[str, Any]:
    dumped = data.model_dump()
    return {
        "header": dumped["header"],
        "shipment": dumped["shipment"],
        "items": dumped["items"],
        "declaration": dumped,
        "today": date.today().isoformat(),
    }


def render_export(data: DeclarationData, template_path: Path | None, base_name: str) -> tuple[bytes, str, str]:
    """Возвращает (содержимое файла, имя файла, MIME-тип)."""
    path = template_path or DEFAULT_TEMPLATE
    ext = path.suffix.lower()
    try:
        if ext in TEXT_TEMPLATE_TYPES:
            return _render_text(path, data), f"{base_name}{ext}", TEXT_TEMPLATE_TYPES[ext]
        if ext in {".xlsx", ".xlsm"}:
            return _render_xlsx(path, data), f"{base_name}.xlsx", XLSX_MEDIA_TYPE
    except TemplateError as exc:
        raise TemplateRenderError(t("export.template_error", error=exc)) from exc
    raise TemplateRenderError(t("export.unsupported", ext=ext))


def _render_text(path: Path, data: DeclarationData) -> bytes:
    # Для XML включаем экранирование: «ООО "Ромашка" & Co» не сломает разметку.
    env = _make_env(autoescape=path.suffix.lower() == ".xml")
    template = env.from_string(path.read_text(encoding="utf-8-sig"))
    return template.render(**_context(data)).encode("utf-8")


def _render_cell(env: SandboxedEnvironment, value: Any, ctx: dict[str, Any]) -> Any:
    if not isinstance(value, str) or "{{" not in value and "{%" not in value:
        return value
    match = _SINGLE_EXPRESSION.match(value)
    if match:
        # Ячейка целиком — одно выражение: сохраняем тип (число останется числом в Excel).
        result = env.compile_expression(match.group(1).strip(), undefined_to_none=True)(**ctx)
        if result is None:
            return None
        if isinstance(result, (int, float, str)):
            return result
        return str(result)
    return env.from_string(value).render(**ctx)


def _render_xlsx(path: Path, data: DeclarationData) -> bytes:
    env = _make_env(autoescape=False)
    ctx = _context(data)
    workbook = load_workbook(path)

    for sheet in workbook.worksheets:
        # 1. Находим строку-образец товарной позиции (первая строка с «item.»).
        item_row = next(
            (row[0].row for row in sheet.iter_rows()
             if any(isinstance(c.value, str) and "item." in c.value for c in row)),
            None,
        )
        if item_row is not None:
            template_cells = [(c.column, c.value, c._style) for c in sheet[item_row] if not isinstance(c, MergedCell)]
            row_height = sheet.row_dimensions[item_row].height
            extra = len(data.items) - 1
            if extra > 0:
                # Ограничение MVP: insert_rows не сдвигает объединённые ячейки и формулы ниже.
                sheet.insert_rows(item_row + 1, extra)
            elif extra < 0:  # позиций нет — удаляем строку-образец
                sheet.delete_rows(item_row)
            for offset, item in enumerate(ctx["items"]):
                row_index = item_row + offset
                sheet.row_dimensions[row_index].height = row_height
                for column, value, style in template_cells:
                    cell = sheet.cell(row=row_index, column=column)
                    cell._style = copy(style)
                    cell.value = _render_cell(env, value, {**ctx, "item": item, "loop_index": offset + 1})

        # 2. Остальные ячейки — значения шапки и итогов.
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):  # значение хранится в левой верхней ячейке
                    continue
                if item_row is not None and item_row <= cell.row < item_row + len(data.items):
                    continue
                cell.value = _render_cell(env, cell.value, ctx)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
