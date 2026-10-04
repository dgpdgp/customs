"""Генерирует демонстрационные Excel-файлы в папке samples/.

    python scripts/make_samples.py

Все данные вымышлены. Файлы уже лежат в репозитории; скрипт нужен, чтобы
их можно было пересоздать или изменить.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SAMPLES = Path(__file__).resolve().parent.parent / "samples"
BOLD = Font(bold=True)
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")

# (артикул, описание, код ТН ВЭД в инвойсе, кол-во, ед., цена, сумма, мест, нетто, брутто)
LINES = [
    ("BRG-6205", "Ball bearing 6205-2RS", None, 600, "pcs", 4.60, 2760.00, 3, 72.0, 78.6),
    ("BLT-SPZ1250", "V-belt SPZ 1250", None, 200, "pcs", 7.50, 1500.00, 2, 28.0, 30.8),
    ("SEAL-35527", "Oil seal 35x52x7 NBR", None, 1000, "pcs", 0.85, 850.00, 1, 9.5, 10.4),
    ("FLT-0450", "Hydraulic filter element 10 micron", "8421230000", 50, "pcs", 12.40, 620.00, 2, 31.0, 34.2),
]


def _table(ws, start_row: int, headers: list[str], rows: list[list]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
    for r, values in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value).border = BOX


def make_invoice() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "COMMERCIAL INVOICE No. INV-2026-118"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Date: 22.09.2026"
    ws["A3"] = "Seller: Beta Maschinenbau GmbH, Industriestraße 12, 70565 Stuttgart, Germany"
    ws["A4"] = "Buyer: Alfa Import LLC (ООО «Альфа Импорт»), Tashkent, Uzbekistan"
    ws["A5"] = "Contract: AB-2025/17 dd 15.03.2025   Terms: DAP Tashkent (Incoterms 2020)   Currency: EUR"
    ws["A6"] = "Country of origin: Germany"
    rows = [[i, a, d, hs, q, u, p, s] for i, (a, d, hs, q, u, p, s, *_rest) in enumerate(LINES, start=1)]
    _table(ws, 8, ["No", "Article", "Description", "HS code", "Qty", "Unit", "Unit price, EUR", "Amount, EUR"], rows)
    total_row = 9 + len(LINES)
    ws.cell(row=total_row, column=7, value="TOTAL, EUR").font = BOLD
    ws.cell(row=total_row, column=8, value=sum(line[6] for line in LINES)).font = BOLD
    for col, width in zip("ABCDEFGH", [5, 14, 36, 12, 8, 6, 15, 13], strict=True):
        ws.column_dimensions[col].width = width
    wb.save(SAMPLES / "invoice_INV-2026-118.xlsx")


def make_packing_list() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Packing list"
    ws["A1"] = "PACKING LIST to invoice INV-2026-118 dd 22.09.2026"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Transport: CMR 0461190, truck S 517 KL / SA 8840"
    rows = [[i, a, d, q, pk, n, g] for i, (a, d, _hs, q, _u, _p, _s, pk, n, g) in enumerate(LINES, start=1)]
    _table(ws, 4, ["No", "Article", "Description", "Qty", "Packages", "Net weight, kg", "Gross weight, kg"], rows)
    total_row = 5 + len(LINES)
    ws.cell(row=total_row, column=4, value="TOTAL").font = BOLD
    ws.cell(row=total_row, column=5, value=sum(line[7] for line in LINES)).font = BOLD
    ws.cell(row=total_row, column=6, value=round(sum(line[8] for line in LINES), 3)).font = BOLD
    ws.cell(row=total_row, column=7, value=round(sum(line[9] for line in LINES), 3)).font = BOLD
    for col, width in zip("ABCDEFG", [5, 14, 36, 8, 10, 15, 16], strict=True):
        ws.column_dimensions[col].width = width
    wb.save(SAMPLES / "packing_list_INV-2026-118.xlsx")


def make_excel_template() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Декларация"
    ws["A1"] = "Таможенная декларация — черновик для ввода"
    ws["A1"].font = Font(bold=True, size=14)
    header = [
        ("Тип декларации", "{{ header.declaration_type }}"),
        ("Таможенный пост", "{{ header.customs_office }}"),
        ("Отправитель", "{{ header.exporter.name }}, {{ header.exporter.address }}"),
        ("Получатель", "{{ header.importer.name }}, ИНН {{ header.importer.tax_id }}"),
        ("Декларант", "{{ header.declarant.name }}, ИНН {{ header.declarant.tax_id }}"),
        ("Контракт", "№ {{ header.contract_number }} от {{ header.contract_date }}"),
        ("Условия поставки", "{{ header.delivery_terms }} {{ header.delivery_place }}"),
        ("Инвойс", "{{ shipment.invoice_numbers|join(', ') }} от {{ shipment.invoice_date }}"),
        ("Транспорт", "{{ shipment.transport_document }}, ТС {{ shipment.vehicle_id }}"),
        ("Валюта", "{{ header.currency }}"),
    ]
    for r, (label, value) in enumerate(header, start=3):
        ws.cell(row=r, column=1, value=label).font = BOLD
        ws.cell(row=r, column=2, value=value)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)

    columns = [
        ("№", "{{ item.item_no }}"), ("Описание", "{{ item.description }}"), ("Артикул", "{{ item.article }}"),
        ("Код ТН ВЭД", "{{ item.hs_code }}"), ("Страна", "{{ item.country_of_origin }}"),
        ("Кол-во", "{{ item.quantity }}"), ("Ед.", "{{ item.unit }}"), ("Мест", "{{ item.packages }}"),
        ("Брутто, кг", "{{ item.gross_weight_kg }}"), ("Нетто, кг", "{{ item.net_weight_kg }}"),
        ("Стоимость", "{{ item.total_value }}"),
    ]
    head_row = 4 + len(header)
    for col, (title, placeholder) in enumerate(columns, start=1):
        cell = ws.cell(row=head_row, column=col, value=title)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
        cell = ws.cell(row=head_row + 1, column=col, value=placeholder)  # строка-образец позиции
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    total_row = head_row + 2
    ws.cell(row=total_row, column=7, value="Итого:").font = BOLD
    for col, expr in [(8, "shipment.total_packages"), (9, "shipment.total_gross_weight_kg"),
                      (10, "shipment.total_net_weight_kg"), (11, "shipment.total_invoice_value")]:
        ws.cell(row=total_row, column=col, value="{{ " + expr + " }}").font = BOLD
    for col, width in zip("ABCDEFGHIJK", [5, 40, 14, 13, 8, 9, 6, 7, 11, 11, 12], strict=True):
        ws.column_dimensions[col].width = width
    wb.save(SAMPLES / "templates" / "declaration_template.xlsx")


if __name__ == "__main__":
    (SAMPLES / "templates").mkdir(parents=True, exist_ok=True)
    make_invoice()
    make_packing_list()
    make_excel_template()
    print("Готово:", *sorted(p.relative_to(SAMPLES) for p in SAMPLES.rglob("*.xlsx")), sep="\n  ")
