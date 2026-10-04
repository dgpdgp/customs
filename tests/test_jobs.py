"""Сквозной сценарий: загрузка -> обработка (фейковая LLM) -> Diff View данные -> утверждение -> экспорт."""

import io

from defusedxml import ElementTree
from openpyxl import load_workbook

from tests.conftest import register_and_login, upload_samples


def _job(client, job_id=1):
    return client.get(f"/api/jobs/{job_id}").json()


def test_upload_requires_login(client):
    response = upload_samples(client)
    assert response.status_code == 303
    assert response.headers["location"].startswith("/login")


def test_full_flow_header_preserved_items_replaced(logged_in, fake_llm):
    response = upload_samples(logged_in)
    assert response.status_code == 303 and response.headers["location"] == "/jobs/1"
    assert fake_llm["documents"] == [
        "reference_declaration.xml", "invoice_INV-2026-118.xlsx", "packing_list_INV-2026-118.xlsx"]

    job = _job(logged_in)
    assert job["status"] == "review", job["error_message"]
    # Реквизиты скопированы из эталона, товары — новые
    assert job["proposed"]["header"] == job["reference"]["header"]
    assert [i["article"] for i in job["proposed"]["items"]] == ["BRG-6205", "BLT-SPZ1250", "SEAL-35527", "FLT-0450"]
    assert job["proposed"]["shipment"]["invoice_numbers"] == ["INV-2026-118"]

    issues = job["issues"]
    # Код не найден -> ошибка (и от LLM, и от валидатора)
    assert {(i["item_no"], i["field"], i["source"]) for i in issues if i["severity"] == "error"} == {
        (3, "hs_code", "llm"), (3, "hs_code", "validator")}
    # Цитаты и числа корректных позиций подтверждены — замечаний к ним нет
    assert not [i for i in issues if i["field"] in ("source_quote", "total_value", "net_weight_kg")]


def test_validator_catches_hallucinated_line(logged_in, fake_llm):
    item = fake_llm["result"].new_items[0]
    item.source_quote = "1 | BRG-9999 | Ball bearing premium | 650 | pcs | 4.6 | 2990"
    item.total_value = 2990.0  # число, которого нет в документах
    upload_samples(logged_in)

    issues = _job(logged_in)["issues"]
    messages = {(i["item_no"], i["field"]): i["message"] for i in issues if i["source"] == "validator"}
    assert "не найдена в документах" in messages[(1, "source_quote")]
    assert "не встречается" in messages[(1, "total_value")]
    assert "не совпадает с итогом" in messages[(None, "shipment.total_invoice_value")]


def test_validator_catches_wrong_reference_code(logged_in, fake_llm):
    fake_llm["result"].new_items[1].hs_code = "4010390000"  # «по эталону», но в эталоне другой код
    upload_samples(logged_in)
    issues = _job(logged_in)["issues"]
    assert any(i["item_no"] == 2 and i["field"] == "hs_code" and "в позиции эталона" in i["message"] for i in issues)


def test_group_by_hs_code(logged_in, fake_llm):
    # Позиции 1 и 2 с одинаковым кодом и страной объединяются; суммы считает сервер
    fake_llm["result"].new_items[1].hs_code = "8482100009"
    upload_samples(logged_in, group_by_hs=True)
    job = _job(logged_in)
    items = job["proposed"]["items"]
    assert [i["article"] for i in items] == ["BRG-6205, BLT-SPZ1250", "SEAL-35527", "FLT-0450"]
    merged = items[0]
    assert merged["quantity"] == 800 and merged["total_value"] == 4260.0
    assert merged["net_weight_kg"] == 100.0 and merged["packages"] == 5
    assert merged["unit_price"] is None
    # Замечание LLM к исходной позиции 3 перенумеровано вслед за позицией (3 -> 2)
    llm_issue = next(i for i in job["issues"] if i["source"] == "llm" and i["field"] == "hs_code")
    assert llm_issue["item_no"] == 2


def test_failed_llm_shows_error_and_retry(logged_in, monkeypatch, fake_llm):
    from app.services import llm

    def boom(*_args):
        raise llm.LLMError("Неверный ключ Anthropic API (ANTHROPIC_API_KEY)")

    monkeypatch.setattr(llm, "extract_declaration", boom)
    upload_samples(logged_in)
    job = _job(logged_in)
    assert job["status"] == "failed" and "ключ" in job["error_message"]

    monkeypatch.setattr(llm, "extract_declaration", lambda r, c: (fake_llm["result"], llm.LLMCallInfo("m", 1, 1, 1)))
    assert logged_in.post("/api/jobs/1/retry").json()["status"] == "processing"
    assert _job(logged_in)["status"] == "review"


def test_upload_rejects_bad_extension(logged_in, fake_llm):
    response = logged_in.post("/jobs", files=[
        ("reference_file", ("ref.exe", b"MZ...", "application/octet-stream")),
        ("commercial_files", ("inv.pdf", b"%PDF", "application/pdf")),
    ])
    assert response.status_code == 400
    assert "формат .exe не подходит" in response.text
    assert fake_llm["calls"] == 0


def test_other_user_cannot_see_job(client, fake_llm):
    register_and_login(client, "owner@example.com")
    upload_samples(client)
    client.cookies.clear()
    register_and_login(client, "intruder@example.com")
    assert client.get("/api/jobs/1").status_code == 404
    assert client.get("/jobs/1/export").status_code == 404


def test_approve_and_export_default_xml(logged_in, fake_llm):
    upload_samples(logged_in)
    assert logged_in.get("/jobs/1/export").status_code == 409  # до утверждения нельзя

    data = _job(logged_in)["proposed"]
    data["items"][2]["hs_code"] = "4016930001"  # декларант вписал код вручную
    result = logged_in.post("/api/jobs/1/approve", json={"data": data}).json()
    assert result["status"] == "approved"
    assert not [i for i in result["issues"] if i["severity"] == "error"]

    response = logged_in.get("/jobs/1/export")
    assert response.status_code == 200
    assert "declaration_1.xml" in response.headers["content-disposition"]
    root = ElementTree.fromstring(response.content)
    assert root.find("Header/Importer/Name").text == "ООО «Альфа Импорт»"
    assert root.find("Items").get("count") == "4"
    assert [e.text for e in root.iter("HSCode")] == ["8482100009", "4010320000", "4016930001", "8421230000"]
    assert _job(logged_in)["status"] == "exported"


def test_approve_rejects_malformed_data(logged_in, fake_llm):
    upload_samples(logged_in)
    data = _job(logged_in)["proposed"]
    data["items"][0]["total_value"] = "две тысячи"
    assert logged_in.post("/api/jobs/1/approve", json={"data": data}).status_code == 422


def test_export_with_asycuda_like_template(logged_in, fake_llm):
    upload_samples(logged_in, template="asycuda_like_template.xml")
    logged_in.post("/api/jobs/1/approve", json={"data": _job(logged_in)["proposed"]})
    root = ElementTree.fromstring(logged_in.get("/jobs/1/export").content)
    assert root.tag == "ASYCUDA"
    assert len(root.findall("Item")) == 4
    assert root.find("Traders/Consignee/Consignee_name").text == "ООО «Альфа Импорт»"
    assert root.find("Valuation/Total_invoice/Amount_foreign_currency").text == "5730.00"


def test_export_with_excel_template(logged_in, fake_llm):
    upload_samples(logged_in, template="declaration_template.xlsx")
    logged_in.post("/api/jobs/1/approve", json={"data": _job(logged_in)["proposed"]})
    response = logged_in.get("/jobs/1/export")
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")

    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = [[c for c in row] for row in sheet.iter_rows(values_only=True)]
    assert rows[5][1] == "ООО «Альфа Импорт», ИНН 300000001"
    item_rows = [r for r in rows if isinstance(r[0], int)]
    assert [r[2] for r in item_rows] == ["BRG-6205", "BLT-SPZ1250", "SEAL-35527", "FLT-0450"]
    assert item_rows[0][10] == 2760.0  # число осталось числом
    total_row = next(r for r in rows if r[6] == "Итого:")
    assert total_row[10] == 5730.0


def test_review_page_renders(logged_in, fake_llm):
    upload_samples(logged_in)
    page = logged_in.get("/jobs/1")
    assert page.status_code == 200
    assert "diff.js" in page.text and "Задача №1" in page.text
